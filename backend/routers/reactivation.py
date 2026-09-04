# backend/routers/reactivation.py
from typing import Optional
from pydantic import BaseModel

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, status
from sqlalchemy.orm import Session
from sqlalchemy import text, or_
import pandas as pd
import io
import os
import json
import shutil
import mimetypes
from fastapi.responses import FileResponse
from datetime import date
from fastapi.responses import StreamingResponse
from backend.utils.exporter import generate_excel_export,generate_csv_export
from backend.database import SessionLocal
from backend.routers.auth import get_current_user
from backend.models import User, District
from backend.models.base import to_name, get_ist_now, StatusEnum


# 🌟 CONNECTED: Pulling from your exact, unchanged model class baseline definitions
from backend.models.reactivation import (
    OperatorReactivationRequest, 
    ReactivationDocument, 
    ReactivationOperator, 
    ReactivationRemarkHistory
)
from backend.models.operator import Operator

router = APIRouter()

BASE_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
UPLOAD_DIR = os.path.join(BASE_DIR, "uploads", "reactivation")
os.makedirs(UPLOAD_DIR, exist_ok=True)

def get_user_role_str(current_user) -> str:
    if hasattr(current_user.role, "role"):
        return str(current_user.role.role).lower()
    elif hasattr(current_user.role, "value"):
        return str(current_user.role.value).lower()
    elif hasattr(current_user.role, "name"):
        return str(current_user.role.name).lower()
    return str(current_user.role).lower()

def validate_attendance_excel_records(attendance_file_or_path, operator_rows):
    """
    Parses the Attendance Excel sheet and validates that each operator in operator_rows
    matches BOTH their name and mobile number within the file.
    """
    import io
    import re
    import os
    import openpyxl
    import pandas as pd
    from backend.utils.ocr_utils import _match_operator_name

    raw_bytes = None
    try:
        if isinstance(attendance_file_or_path, bytes):
            raw_bytes = attendance_file_or_path
        elif hasattr(attendance_file_or_path, 'file') and hasattr(attendance_file_or_path.file, 'read'):
            attendance_file_or_path.file.seek(0)
            raw_bytes = attendance_file_or_path.file.read()
            attendance_file_or_path.file.seek(0)
        elif hasattr(attendance_file_or_path, 'read') and not hasattr(attendance_file_or_path, '__await__'):
            raw_bytes = attendance_file_or_path.read()
            if hasattr(attendance_file_or_path, 'seek'):
                attendance_file_or_path.seek(0)
        elif isinstance(attendance_file_or_path, str):
            resolved_path = attendance_file_or_path
            if not os.path.exists(resolved_path) and not os.path.isabs(resolved_path):
                candidate_p = os.path.join(BASE_DIR, resolved_path.lstrip("/\\"))
                if os.path.exists(candidate_p):
                    resolved_path = candidate_p
            if os.path.exists(resolved_path):
                with open(resolved_path, 'rb') as f:
                    raw_bytes = f.read()
    except Exception as e:
        print(f"Warning reading attendance file bytes: {e}")

    if not raw_bytes:
        return [f"'{op.get('name', op.get('operator_name', 'Operator'))}' (Empty or unreadable Excel sheet)" for op in operator_rows]

    rows = []
    # 1. Try openpyxl across all worksheets (preserves exact values, numbers, formats)
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True)
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cell_strs = []
                for cell in row:
                    if cell is None:
                        continue
                    if isinstance(cell, float) and cell.is_integer():
                        cell_strs.append(str(int(cell)))
                    elif isinstance(cell, (int, float)):
                        cell_strs.append(f"{cell:.0f}" if isinstance(cell, float) and cell.is_integer() else str(cell))
                    else:
                        cell_strs.append(str(cell).strip())
                cell_strs = [c for c in cell_strs if c and c.lower() not in ("none", "nan", "")]
                if cell_strs:
                    rows.append(" ".join(cell_strs))
    except Exception:
        pass

    # 2. Fallback to pandas with header=None across all sheets
    if not rows:
        try:
            excel_file = pd.ExcelFile(io.BytesIO(raw_bytes))
            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(excel_file, sheet_name=sheet_name, header=None, dtype=str)
                for _, r in df.iterrows():
                    cell_strs = [str(val).strip() for val in r.values if pd.notna(val) and str(val).strip().lower() not in ("nan", "none", "")]
                    if cell_strs:
                        rows.append(" ".join(cell_strs))
        except Exception:
            pass

    # 3. Fallback to CSV
    if not rows:
        try:
            for enc in ['utf-8', 'latin-1', 'cp1252']:
                try:
                    df = pd.read_csv(io.BytesIO(raw_bytes), header=None, dtype=str, encoding=enc)
                    for _, r in df.iterrows():
                        cell_strs = [str(val).strip() for val in r.values if pd.notna(val) and str(val).strip().lower() not in ("nan", "none", "")]
                        if cell_strs:
                            rows.append(" ".join(cell_strs))
                    if rows:
                        break
                except Exception:
                    pass
        except Exception:
            pass

    if not rows:
        return [f"'{op.get('name', op.get('operator_name', 'Operator'))}' (Empty or unreadable Excel sheet)" for op in operator_rows]

    from backend.utils.ocr_utils import _match_operator_in_excel_row

    missing = []
    for op in operator_rows:
        name = str(op.get('name', op.get('operator_name', ''))).strip()
        mobile = str(op.get('mobile', op.get('operator_mobile', ''))).strip()

        row_matched = False
        found_mob_anywhere = False
        found_name_anywhere = False

        for row_text in rows:
            mob_in_row, name_in_row = _match_operator_in_excel_row(name, mobile, row_text)
            if mob_in_row:
                found_mob_anywhere = True
            if name_in_row:
                found_name_anywhere = True

            # STRICT COMBINATION MATCH: Both must match in the exact same row
            if mob_in_row and name_in_row:
                row_matched = True
                break

        if row_matched:
            continue

        if found_mob_anywhere and not found_name_anywhere:
            missing.append(f"'{name or 'Operator'}' (Mobile {mobile} matched, but name did not match row)")
        elif found_name_anywhere and not found_mob_anywhere:
            missing.append(f"'{name or 'Operator'}' (Name matched, but mobile {mobile} not found in row)")
        elif found_mob_anywhere and found_name_anywhere:
            missing.append(f"'{name or 'Operator'}' (Mobile and name found in different rows, must be in same row)")
        else:
            missing.append(f"'{name or 'Operator'}' (Mobile: {mobile or 'N/A'})")

    return missing

def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()

def generate_dynamic_request_code(db: Session, district_id: str, district_name: str) -> str:
    district = db.query(District).filter(District.district_code == str(district_id)).first()
    prefix = district.district_short_name if district else None
    
    if not prefix:
        name_clean = district_name.strip().lower()
        district_map = {"raipur": "RP", "bilaspur": "BP", "durg": "DG"}
        prefix = district_map.get(name_clean, "".join([c for c in name_clean if c.isalnum()])[:2].upper())
        if len(prefix) != 2: prefix = "XX"

    
    last_req = db.query(OperatorReactivationRequest).filter(
        OperatorReactivationRequest.district_id == district_id
    ).order_by(OperatorReactivationRequest.id.desc()).first()
    
    next_num = 1
    if last_req and last_req.request_code and "-R" in last_req.request_code:
        try:
            next_num = int(last_req.request_code.split("-R")[1]) + 1
        except ValueError:
            next_num = db.query(OperatorReactivationRequest).filter(
                OperatorReactivationRequest.district_id == district_id
            ).count() + 1

    else:
        next_num = db.query(OperatorReactivationRequest).filter(
            OperatorReactivationRequest.district_id == district_id
        ).count() + 1

    return f"{prefix}-R{next_num:04d}"


@router.post("/submit", status_code=status.HTTP_201_CREATED)
async def submit_operator_reactivation(
    training_date: str = Form(...),
    training_photo: Optional[UploadFile] = File(None),
    nodal_letter: Optional[UploadFile] = File(None),
    om_letter: Optional[UploadFile] = File(None),
    attendance_list: Optional[UploadFile] = File(None),
    manual_operators: str = Form(...), 
    reapply_request_code: str = Form(None),
    dc_remark: str = Form(None),

    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    try:
        user_role_str = get_user_role_str(current_user)
        if user_role_str not in ["dc", "edm"]: 
            raise HTTPException(status_code=403, detail="Access Denied. Only DC/EDM can submit requests.")
        if not current_user.district_id: 
            raise HTTPException(status_code=400, detail="Missing user district layout configuration mapping.")

        try:
            operator_rows = json.loads(manual_operators)
        except (json.JSONDecodeError, TypeError):
            raise HTTPException(status_code=400, detail="Invalid operator rows payload JSON data structure matrix.")

        if len(operator_rows) == 0:
            raise HTTPException(status_code=400, detail="The operator log datagrid cannot be processed with zero rows.")

        # 📊 Validate Attendance Excel sheet contents against submitted operator rows
        has_new_attendance = False
        if attendance_list and getattr(attendance_list, 'filename', None):
            try:
                attendance_list.file.seek(0, 2)
                if attendance_list.file.tell() > 0:
                    has_new_attendance = True
                attendance_list.file.seek(0)
            except Exception:
                has_new_attendance = True

        target_attendance_input = attendance_list if has_new_attendance else None
        if not target_attendance_input and reapply_request_code:
            existing_doc = db.query(ReactivationDocument).join(
                OperatorReactivationRequest, ReactivationDocument.request_id == OperatorReactivationRequest.id
            ).filter(
                OperatorReactivationRequest.request_code == reapply_request_code,
                ReactivationDocument.doc_type == "attendance_list"
            ).first()
            if existing_doc and existing_doc.path and os.path.exists(existing_doc.path):
                target_attendance_input = existing_doc.path

        if target_attendance_input:
            missing_ops = validate_attendance_excel_records(target_attendance_input, operator_rows)
            if missing_ops:
                missing_str = ", ".join(missing_ops)
                raise HTTPException(
                    status_code=400,
                    detail=f"Operator Attendance Excel Sheet Validation Failed: The uploaded Excel does not contain details for: {missing_str}. Please upload an Attendance Excel Sheet containing all operators added to this batch."
                )

        district = db.query(District).filter(District.district_code == str(current_user.district_id)).first()
        district_name = district.district_name if district else "Unknown"
        district_folder = district_name.strip().replace(" ", "_").upper()

        if reapply_request_code:
            req = db.query(OperatorReactivationRequest).filter(
                OperatorReactivationRequest.request_code == reapply_request_code,
                OperatorReactivationRequest.dc_id == current_user.id
            ).first()
            if not req:
                raise HTTPException(status_code=404, detail="Original request not found or access denied.")
            
            req_code = reapply_request_code
            
            req.training_date = date.fromisoformat(training_date.strip())
            req.status_id = StatusEnum.REAPPLIED.value
            req.is_mailed = 0
        else:
            req_code = generate_dynamic_request_code(db, str(current_user.district_id), district_name)
    
            new_request = OperatorReactivationRequest(
                request_code=req_code,
                dc_id=current_user.id,
                district_id=str(current_user.district_id),
                operator_count=len(operator_rows),
                training_date=date.fromisoformat(training_date.strip()),
                status="PENDING"
            )
            db.add(new_request)
            db.flush()  # To get new_request.id

        active_req_id = req.id if reapply_request_code else new_request.id
        request_folder = os.path.join(BASE_DIR, "uploads", "reactivation", district_folder, req_code)

        os.makedirs(request_folder, exist_ok=True)
        document_files = {"training_photo": training_photo, "nodal_letter": nodal_letter, "om_letter": om_letter, "attendance_list": attendance_list}

        for doc_type, uploaded_file in document_files.items():
            if not uploaded_file or uploaded_file.filename == '':
                continue  # No new file uploaded, retain the existing one
                
            # If replacing an existing document during reapply, delete the old record for this doc_type
            if reapply_request_code:
                db.query(ReactivationDocument).filter(ReactivationDocument.request_id == active_req_id, ReactivationDocument.doc_type == doc_type).delete()

            uploaded_file.file.seek(0, 2)
            bytes_size = uploaded_file.file.tell()
            uploaded_file.file.seek(0)
            file_save_path = os.path.join(request_folder, f"{doc_type}_{uploaded_file.filename}")
            with open(file_save_path, "wb") as buffer:
                shutil.copyfileobj(uploaded_file.file, buffer)
            db.add(ReactivationDocument(request_id=active_req_id, doc_type=doc_type, path=file_save_path, original_filename=uploaded_file.filename, file_size=bytes_size))

        from datetime import timedelta
        operators_added = []
        for op in operator_rows:
            op_mobile = str(op.get('mobile', '')).strip()
            op_email = str(op.get('email', '')).strip()
            op_id = op.get('id')
            
            # Check for existing duplicate in DB outside this batch/operator
            mobile_query = db.query(ReactivationOperator).filter(ReactivationOperator.operator_mobile == op_mobile)
            if op_id:
                mobile_query = mobile_query.filter(ReactivationOperator.id != int(op_id))
            elif reapply_request_code:
                mobile_query = mobile_query.filter(ReactivationOperator.request_id != req.id)
            if mobile_query.first():
                raise HTTPException(status_code=400, detail=f"Operator Reactivation Request already exists with mobile number: {op_mobile}")
                
            if op_email:
                email_query = db.query(ReactivationOperator).filter(ReactivationOperator.email_id == op_email)
                if op_id:
                    email_query = email_query.filter(ReactivationOperator.id != int(op_id))
                elif reapply_request_code:
                    email_query = email_query.filter(ReactivationOperator.request_id != req.id)
                if email_query.first():
                    raise HTTPException(status_code=400, detail=f"Operator Reactivation Request already exists with email address: {op_email}")

            parsed_cert_date = date.fromisoformat(op['certDate']) if op.get('certDate') else None
            if parsed_cert_date:
                three_years_ago = date.today() - timedelta(days=3*365)
                if parsed_cert_date < three_years_ago:
                    raise HTTPException(status_code=400, detail=f"NSEIT Certification Date for {op.get('name')} must not be more than 3 years old.")

            existing_op = None
            if op_id:
                try:
                    existing_op = db.query(ReactivationOperator).filter(
                        ReactivationOperator.id == int(op_id),
                        ReactivationOperator.request_id == active_req_id
                    ).first()
                except Exception:
                    existing_op = None
                    
            if not existing_op and reapply_request_code:
                existing_op = db.query(ReactivationOperator).filter(
                    ReactivationOperator.request_id == active_req_id,
                    ReactivationOperator.operator_mobile == op_mobile
                ).first()

            if existing_op:
                existing_op.role = op.get('role', '').strip()
                existing_op.operator_name = str(op.get('name', '')).strip()
                existing_op.registrar_code = op.get('reg', '').strip()
                existing_op.ea_code = op.get('ea', '').strip()
                existing_op.user_code = op.get('user', '').strip()
                existing_op.certificate_number = op.get('cert', '').strip()
                existing_op.lms_certificate_id = op.get('lmsId', '').strip()
                existing_op.operator_mobile = op_mobile
                existing_op.email_id = op_email
                existing_op.aadhaar_number = str(op.get('aadhar', '')).strip()
                existing_op.certification_date = parsed_cert_date
                existing_op.remarks = op.get('remarks', '').strip()
                existing_op.model_type = op.get('model', '').strip()
                existing_op.status_id = StatusEnum.REAPPLIED.value
                existing_op.reject_reason = None
                operators_added.append(existing_op)
            else:
                new_op = ReactivationOperator(
                    request_id=active_req_id,
                    role=op.get('role', '').strip(),
                    operator_name=str(op.get('name', '')).strip(),
                    registrar_code=op.get('reg', '').strip(),
                    ea_code=op.get('ea', '').strip(),
                    user_code=op.get('user', '').strip(),
                    certificate_number=op.get('cert', '').strip(),
                    lms_certificate_id=op.get('lmsId', '').strip(),
                    operator_mobile=op_mobile,
                    email_id=op_email,
                    aadhaar_number=str(op.get('aadhar', '')).strip(),
                    certification_date=parsed_cert_date,
                    remarks=op.get('remarks', '').strip(),
                    model_type=op.get('model', '').strip(),      
                    status_id=StatusEnum.REAPPLIED.value if reapply_request_code else StatusEnum.PENDING.value
                )
                db.add(new_op)
                operators_added.append(new_op)
            
        db.flush()

        if reapply_request_code:
            req.operator_count = db.query(ReactivationOperator).filter(ReactivationOperator.request_id == req.id).count()
        
        for op_item in operators_added:
            msg = f"Operator '{op_item.operator_name}' {'is reapplied' if reapply_request_code else 'is submitted'} by DC"
            if reapply_request_code and dc_remark:
                msg += f". Remarks: {dc_remark.strip()}"
                
            db.add(ReactivationRemarkHistory(
                request_id=active_req_id,
                operator_id=op_item.id,
                author_id=current_user.id,
                remark_history=msg,
                sender_role="DC",
                status_after="REAPPLIED" if reapply_request_code else "PENDING"
            ))
            
        db.commit()
        return {"success": True, "request_code": req_code}
    except Exception as e:
        import traceback
        error_msg = f"FastAPI internal error: {str(e)}\n{traceback.format_exc()}"
        print(error_msg)
        raise HTTPException(status_code=400, detail=str(e))


@router.get("/check-duplicate")
def check_duplicate(
    mobile: str = None, 
    email: str = None, 
    exclude_id: str = None,
    db: Session = Depends(get_db)
):
    parsed_exclude_id = None
    if exclude_id and exclude_id.strip():
        try:
            parsed_exclude_id = int(exclude_id)
        except ValueError:
            pass

    if mobile:
        query = db.query(ReactivationOperator).filter(ReactivationOperator.operator_mobile == mobile.strip())
        if parsed_exclude_id is not None:
            query = query.filter(ReactivationOperator.id != parsed_exclude_id)
        if query.first():
            return {"exists": True, "message": "An Operator Reactivation Request already exists with this mobile number."}
            
    if email:
        query = db.query(ReactivationOperator).filter(ReactivationOperator.email_id == email.strip())
        if parsed_exclude_id is not None:
            query = query.filter(ReactivationOperator.id != parsed_exclude_id)
        if query.first():
            return {"exists": True, "message": "An Operator Reactivation Request already exists with this email address."}
            
    return {"exists": False}

@router.get("/requests")
async def get_reactivation_requests(
    current_user: User = Depends(get_current_user), 
    db: Session = Depends(get_db)
):
    user_role_str = get_user_role_str(current_user)
    
    query = db.query(OperatorReactivationRequest, District.district_name)\
              .outerjoin(District, OperatorReactivationRequest.district_id == District.district_code)
    if user_role_str in ["dc", "edm"]:
        query = query.filter(
            (OperatorReactivationRequest.district_id == str(current_user.district_id)) |
            (OperatorReactivationRequest.dc_id == current_user.id)
        )
    
    requests = query.order_by(OperatorReactivationRequest.updated_at.desc()).all()
        
    compiled_list = []
    for req, dist_name in requests:
        compiled_list.append({
            "id": req.id,
            "request_code": req.request_code,
            "operator_count": req.operator_count,
            "training_date": str(req.training_date) if req.training_date else "",
            "status": to_name(req.status_id).upper().replace(" ", "_").strip(),
            "created_at": str(req.created_at)[:19] if req.created_at else "",
            "updated_at": str(req.updated_at)[:19] if req.updated_at else "",
            "district_name": dist_name or "Raipur",
            "revert_reason": next((r.remark_history for r in reversed(req.remarks) if r.status_after_id in [StatusEnum.REVERTED.value, StatusEnum.REJECTED.value]), "")

        })
    return compiled_list


@router.get("/requests-with-operators")
async def get_reactivation_requests_with_operators(
    current_user: User = Depends(get_current_user), 
    db: Session = Depends(get_db)
):
    user_role_str = get_user_role_str(current_user)
    
    query = db.query(OperatorReactivationRequest, District.district_name)\
              .outerjoin(District, OperatorReactivationRequest.district_id == District.district_code)
    if user_role_str in ["dc", "edm"]:
        query = query.filter(
            (OperatorReactivationRequest.district_id == str(current_user.district_id)) |
            (OperatorReactivationRequest.dc_id == current_user.id)
        )
    
    requests = query.order_by(OperatorReactivationRequest.updated_at.desc()).all()
        
    compiled_list = []
    for req, dist_name in requests:
        operators = db.query(ReactivationOperator).filter(ReactivationOperator.request_id == req.id).all()

        ops_data = [
            {
                "id": op.id,
                "operator_name": op.operator_name,
                "operator_mobile": op.operator_mobile,
                "status": str(op.status).upper().replace(" ", "_"),
                "role": op.role or "Operator",
                "email_id": op.email_id or "",
                "registrar_code": op.registrar_code or "986",
                "ea_code": op.ea_code or "",
                "user_code": op.user_code or "",
                "model_type": op.model_type or "",
                "lms_certificate_id": op.lms_certificate_id or "",
                "certificate_number": op.certificate_number or "",
                "aadhaar_number": op.aadhaar_number or "",

                "certification_date": str(op.certification_date) if op.certification_date else "",
                "remarks": op.remarks or "",
                "reject_reason": op.reject_reason or ""
            }
            for op in operators
        ]
        
        # 🌟 CONNECTED: Order matching your model 'timestamp' attribute mapping definitions
        remarks_hist = db.query(ReactivationRemarkHistory).filter(ReactivationRemarkHistory.request_id == req.id).order_by(ReactivationRemarkHistory.timestamp.asc()).all()

        timeline_logs = [
            {
                "id": rm.id,
                "message": rm.remark_history,
                "sender_role": rm.sender_role,
                "sender_username": rm.author.username if rm.author else "",
                "timestamp": str(rm.timestamp)[:19] if rm.timestamp else "",
                "status_after": rm.status_after,
                "operator_id": rm.operator_id,

            }
            for rm in remarks_hist
        ]
        
        compiled_list.append({
            "id": req.id,
            "request_code": req.request_code,
            "operator_count": req.operator_count,
            "training_date": str(req.training_date) if req.training_date else "",
            "status": to_name(req.status_id).upper().replace(" ", "_").strip(),
            "created_at": str(req.created_at)[:19] if req.created_at else "",
            "updated_at": str(req.updated_at)[:19] if req.updated_at else "",
            "is_mailed": int(req.is_mailed or 0),
            "district_name": dist_name or "Raipur",
            "revert_reason": next((r.remark_history for r in reversed(req.remarks) if r.status_after_id in [StatusEnum.REVERTED.value, StatusEnum.REJECTED.value]), ""),
            "operators": ops_data,
            "timeline_logs": timeline_logs
        })
    return compiled_list


@router.get("/operators/{request_code}")
async def get_individual_operators_by_batch(request_code: str, db: Session = Depends(get_db)):
    req = db.query(OperatorReactivationRequest).filter(OperatorReactivationRequest.request_code == request_code).first()
    if not req:
        raise HTTPException(status_code=404, detail="Request not found")
        
    operators = db.query(ReactivationOperator).filter(ReactivationOperator.request_id == req.id).all()

    ops_data = [
        {
            "id": op.id,
            "operator_name": op.operator_name,
            "operator_mobile": op.operator_mobile,
            "status": str(op.status).upper().replace(" ", "_"),
            "role": op.role or "Operator",
            "email_id": op.email_id or "",
            "registrar_code": op.registrar_code or "986",
            "ea_code": op.ea_code or "",
            "user_code": op.user_code or "",
            "model_type": op.model_type or "",
            "lms_certificate_id": op.lms_certificate_id or "",
            "certificate_number": op.certificate_number or "",
            "aadhaar_number": op.aadhaar_number or "",

            "certification_date": str(op.certification_date) if op.certification_date else "",
            "remarks": op.remarks or "",
            "reject_reason": op.reject_reason or ""
        }
        for op in operators
    ]
    req = db.query(OperatorReactivationRequest).filter(OperatorReactivationRequest.request_code == request_code).first()
    if not req:
        raise HTTPException(status_code=404, detail="Request not found")
        
    remarks_hist = db.query(ReactivationRemarkHistory).filter(ReactivationRemarkHistory.request_id == req.id).order_by(ReactivationRemarkHistory.timestamp.asc()).all()

    timeline_logs = [
        {
            "id": rm.id,
            "message": rm.remark_history,
            "sender_role": rm.sender_role,
            "sender_username": rm.author.username if rm.author else "",
            "timestamp": str(rm.timestamp)[:19] if rm.timestamp else "",
            "status_after": rm.status_after,
            "operator_id": rm.operator_id,

        }
        for rm in remarks_hist
    ]
    
    docs = db.query(ReactivationDocument).filter(ReactivationDocument.request_id == req.id).all()
    docs_data = {
        doc.doc_type: {
            "original_filename": doc.original_filename,
            "path": doc.path
        }
        for doc in docs
    }
    
    # Get batch revert/rejection reason
    latest_revert_log = db.query(ReactivationRemarkHistory).filter(
        ReactivationRemarkHistory.request_id == req.id,
        ReactivationRemarkHistory.status_after_id.in_([StatusEnum.REVERTED.value, StatusEnum.REJECTED.value])
    ).order_by(ReactivationRemarkHistory.timestamp.desc()).first()
    batch_revert_reason = latest_revert_log.remark_history if latest_revert_log else ""

    return {
        "operators": ops_data,
        "timeline_logs": timeline_logs,
        "documents": docs_data,
        "batch_status": req.status,
        "batch_revert_reason": batch_revert_reason

    }


@router.post("/operator/{operator_id}/activate")
async def activate_individual_operator(operator_id: int, reason: Optional[str] = Form(None), db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    op = db.query(ReactivationOperator).filter(ReactivationOperator.id == operator_id).first()
    if op: 
        op.status_id = StatusEnum.APPROVED.value
        if reason:
            op.reject_reason = reason  # Optional: store activate remarks in reject_reason or remarks
        
        # Update parent's updated_at
        parent = db.query(OperatorReactivationRequest).filter(OperatorReactivationRequest.id == op.request_id).first()
        if parent:
            parent.updated_at = get_ist_now()

        remark_text = reason.strip() if reason and reason.strip() else "Approved"
        db.add(ReactivationRemarkHistory(
            request_id=op.request_id,
            operator_id=op.id,
            author_id=current_user.id,
            remark_history=remark_text,
            sender_role="CHIPS_ADMIN",
            status_after_id=op.status_id

        ))
        db.commit()
    return {"success": True}


@router.post("/operator/{operator_id}/send-to-uidai")
async def send_to_uidai_individual_operator(operator_id: int, remarks: Optional[str] = Form(None), db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    op = db.query(ReactivationOperator).filter(ReactivationOperator.id == operator_id).first()
    if op: 
        op.status_id = StatusEnum.SENT_TO_UIDAI.value
        
        # Update parent's updated_at
        parent = db.query(OperatorReactivationRequest).filter(OperatorReactivationRequest.id == op.request_id).first()
        if parent:
            parent.updated_at = get_ist_now()

        remark_text = remarks.strip() if remarks and remarks.strip() else "Sent to UIDAI"
        db.add(ReactivationRemarkHistory(
            request_id=op.request_id,
            operator_id=op.id,
            author_id=current_user.id,
            remark_history=remark_text,
            sender_role="CHIPS_ADMIN",
            status_after_id=op.status_id

        ))
        db.commit()
    return {"success": True}

@router.post("/operator/{operator_id}/revert")
async def revert_individual_operator(operator_id: int, reason: str = Form(...), db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    op = db.query(ReactivationOperator).filter(ReactivationOperator.id == operator_id).first()
    if op: 
        op.status_id = StatusEnum.REVERTED.value
        op.reject_reason = reason
        
        # Update parent's updated_at and reset is_mailed
        parent = db.query(OperatorReactivationRequest).filter(OperatorReactivationRequest.id == op.request_id).first()
        if parent:
            parent.updated_at = get_ist_now()
            parent.is_mailed = 0

        db.add(ReactivationRemarkHistory(
            request_id=op.request_id,
            operator_id=op.id,
            author_id=current_user.id,
            remark_history=reason.strip(),
            sender_role="CHIPS_ADMIN",
            status_after_id=op.status_id
        ))
        db.commit()
    return {"success": True}

@router.post("/operator/{operator_id}/reject")
async def reject_individual_operator(operator_id: int, reason: str = Form(...), db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    op = db.query(ReactivationOperator).filter(ReactivationOperator.id == operator_id).first()
    if op: 
        op.status_id = StatusEnum.REJECTED.value
        op.reject_reason = reason
        
        # Update parent's updated_at and reset is_mailed
        parent = db.query(OperatorReactivationRequest).filter(OperatorReactivationRequest.id == op.request_id).first()
        if parent:
            parent.updated_at = get_ist_now()
            parent.is_mailed = 0

        db.add(ReactivationRemarkHistory(
            request_id=op.request_id,
            operator_id=op.id,
            author_id=current_user.id,
            remark_history=reason.strip(),
            sender_role="CHIPS_ADMIN",
            status_after_id=op.status_id
        ))
        db.commit()
    return {"success": True}


@router.post("/operator/{operator_id}/update_reapply")
async def update_and_reapply_operator(
    operator_id: int,
    operator_name: str = Form(None),
    operator_mobile: str = Form(None),
    email_id: str = Form(None),
    model_type: str = Form(None),
    lms_certificate_id: str = Form(None),
    certificate_number: str = Form(None),
    aadhaar_number: str = Form(None),
    role: str = Form(None),
    registrar_code: str = Form(None),
    ea_code: str = Form(None),
    user_code: str = Form(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    op = db.query(ReactivationOperator).filter(ReactivationOperator.id == operator_id).first()
    if not op:
        raise HTTPException(status_code=404, detail="Operator not found")
    
    if operator_name: op.operator_name = operator_name
    if operator_mobile: op.operator_mobile = operator_mobile
    if email_id: op.email_id = email_id
    if model_type: op.model_type = model_type
    if lms_certificate_id: op.lms_certificate_id = lms_certificate_id
    if certificate_number: op.certificate_number = certificate_number
    if aadhaar_number: op.aadhaar_number = aadhaar_number
    if role: op.role = role
    if registrar_code: op.registrar_code = registrar_code
    if ea_code: op.ea_code = ea_code
    if user_code: op.user_code = user_code
    
    op.status_id = StatusEnum.REAPPLIED.value
    
    # Also update the batch status back to REAPPLIED and reset is_mailed
    req = db.query(OperatorReactivationRequest).filter(OperatorReactivationRequest.id == op.request_id).first()
    if req:
        req.updated_at = get_ist_now()
        req.is_mailed = 0
        if req.status_id in [StatusEnum.REVERTED.value, StatusEnum.REJECTED.value, StatusEnum.REVERTED_BY_CHIPS.value]:
            req.status_id = StatusEnum.REAPPLIED.value

    db.add(ReactivationRemarkHistory(
        request_id=req.id,
        operator_id=op.id,
        author_id=current_user.id,
        remark_history=f"Operator '{op.operator_name}' Reapplied and Details Updated",
        sender_role="DC",
        status_after_id=req.status_id if req else StatusEnum.REAPPLIED.value
    ))

    db.commit()
    return {"success": True}


@router.post("/requests/{request_code}/finalize")
async def finalize_batch_request(request_code: str, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):

    try:
        req = db.query(OperatorReactivationRequest).filter(OperatorReactivationRequest.request_code == request_code).first()
        if req: 
            operators = db.query(ReactivationOperator).filter(
                ReactivationOperator.request_id == req.id
            ).all()
            statuses = [op.status_id for op in operators]
            
            if req.status_id in [StatusEnum.PENDING.value, StatusEnum.REAPPLIED.value]:
                if StatusEnum.SENT_TO_UIDAI.value in statuses:
                    req.status_id = StatusEnum.SENT_TO_UIDAI.value
                    req.reviewed_by = current_user.id
                else:
                    req.status_id = StatusEnum.REVERTED.value
                    req.reviewed_by = current_user.id
            elif req.status_id == StatusEnum.SENT_TO_UIDAI.value:
                req.status_id = StatusEnum.APPROVED.value
                req.reviewed_by = current_user.id
            else:
                if StatusEnum.REVERTED.value in statuses and StatusEnum.APPROVED.value not in statuses:
                    req.status_id = StatusEnum.REVERTED.value
                    req.reviewed_by = current_user.id
                else:
                    req.status_id = StatusEnum.APPROVED.value
                    req.reviewed_by = current_user.id
                    
            db.add(ReactivationRemarkHistory(
                request_id=req.id,
                author_id=current_user.id, 
                remark_history=f"Batch Finalized. New Status: {req.status}", 
                sender_role="CHIPS_ADMIN", 
                status_after_id=req.status_id
            ))

            db.commit()
        return {"success": True}
    except Exception as e:
        import traceback
        error_msg = f"FastAPI internal error: {str(e)}\n{traceback.format_exc()}"
        print(error_msg)
        return {"success": False, "error": str(e)}


@router.post("/requests/{request_code}/revert")
async def revert_batch_request(request_code: str, revert_reason: str = Form(...), db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    req = db.query(OperatorReactivationRequest).filter(OperatorReactivationRequest.request_code == request_code).first()
    if req: 
        req.status_id = StatusEnum.REVERTED.value
        req.is_mailed = 0
        req.reviewed_by = current_user.id
        req.updated_at = get_ist_now()
        
        operators = db.query(ReactivationOperator).filter(ReactivationOperator.request_id == req.id).all()
        for op in operators:
            op.status_id = StatusEnum.REVERTED.value
            op.reject_reason = revert_reason
            db.add(ReactivationRemarkHistory(
                request_id=req.id,
                operator_id=op.id,
                author_id=current_user.id,
                remark_history=revert_reason.strip(),
                sender_role="CHIPS",
                status_after_id=StatusEnum.REVERTED.value
            ))
            

        db.commit()
    return {"success": True}


@router.post("/requests/{request_code}/send-to-uidai")
async def backend_batch_request_to_uidai(request_code: str, remarks: str = Form(...), db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    req = db.query(OperatorReactivationRequest).filter(OperatorReactivationRequest.request_code == request_code).first()
    if req: 
        req.status_id = StatusEnum.SENT_TO_UIDAI.value
        req.reviewed_by = current_user.id
        req.updated_at = get_ist_now()
        
        operators = db.query(ReactivationOperator).filter(
            ReactivationOperator.request_id == req.id,
            ReactivationOperator.status_id.in_([StatusEnum.PENDING.value, StatusEnum.REAPPLIED.value])
        ).all()
        for op in operators:
            op.status_id = StatusEnum.SENT_TO_UIDAI.value
            remark_text = remarks.strip() if remarks and remarks.strip() else "Sent to UIDAI"
            db.add(ReactivationRemarkHistory(
                request_id=req.id,
                operator_id=op.id,
                author_id=current_user.id,
                remark_history=remark_text,
                sender_role="CHIPS",
                status_after_id=StatusEnum.SENT_TO_UIDAI.value
            ))
            
        db.commit()
    return {"success": True}


@router.post("/requests/{request_code}/approve-all")
async def approve_all_operators_in_batch(request_code: str, reason: Optional[str] = Form(None), db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    req = db.query(OperatorReactivationRequest).filter(OperatorReactivationRequest.request_code == request_code).first()
    if req:
        req.status_id = StatusEnum.APPROVED.value
        req.reviewed_by = current_user.id
        req.updated_at = get_ist_now()
        
        operators = db.query(ReactivationOperator).filter(ReactivationOperator.request_id == req.id).all()
        for op in operators:
            op.status_id = StatusEnum.APPROVED.value
            if reason:
                op.reject_reason = reason
            
            remark_text = reason.strip() if reason and reason.strip() else "Approved"
            db.add(ReactivationRemarkHistory(
                request_id=req.id,
                operator_id=op.id,
                author_id=current_user.id,
                remark_history=remark_text,
                sender_role="CHIPS_ADMIN",
                status_after_id=StatusEnum.APPROVED.value
            ))
        
        db.commit()
    return {"success": True}


@router.post("/requests/{request_code}/reject-all")
async def reject_all_operators_in_batch(request_code: str, reason: str = Form(...), db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    req = db.query(OperatorReactivationRequest).filter(OperatorReactivationRequest.request_code == request_code).first()
    if req:
        req.status_id = StatusEnum.APPROVED.value
        req.reviewed_by = current_user.id
        req.updated_at = get_ist_now()
        
        operators = db.query(ReactivationOperator).filter(ReactivationOperator.request_id == req.id).all()
        for op in operators:
            op.status_id = StatusEnum.REJECTED.value
            op.reject_reason = reason
            
            db.add(ReactivationRemarkHistory(
                request_id=req.id,
                operator_id=op.id,
                author_id=current_user.id,
                remark_history=reason.strip(),
                sender_role="CHIPS_ADMIN",
                status_after_id=StatusEnum.REJECTED.value
            ))
        
        db.commit()

    return {"success": True}


@router.get("/export-excel/{request_code}")
def export_operators_to_excel_stream(request_code: str, db: Session = Depends(get_db)):
    operators = db.query(ReactivationOperator).join(
        OperatorReactivationRequest, ReactivationOperator.request_id == OperatorReactivationRequest.id
    ).filter(
        OperatorReactivationRequest.request_code == request_code,
        ReactivationOperator.status_id != StatusEnum.REJECTED.value
    ).all()
    
    def get_display_status(status_str):
        s = str(status_str or '').upper().strip()
        return "APPROVED" if s in ["ACTIVE", "APPROVED"] else s

    df = pd.DataFrame([{
        "S.No": i + 1,
        "Role": o.role,
        "Name As Per Aadhaar": o.operator_name,
        "Aadhaar Number": o.aadhaar_number,

        "Registrar Code": o.registrar_code,
        "EA Code": o.ea_code,
        "User Code": o.user_code,
        "NSEIT Certificate Number": o.certificate_number,
        "LMS Certificate ID": o.lms_certificate_id,
        "Mobile": o.operator_mobile,
        "Primary E-MAIL ID": o.email_id,
        "Status": get_display_status(o.status)

    } for i, o in enumerate(operators)])
    excel_buffer = io.BytesIO()
    with pd.ExcelWriter(excel_buffer, engine='openpyxl') as writer: 
        df.to_excel(writer, index=False)
    excel_buffer.seek(0)
    return StreamingResponse(excel_buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f"attachment; filename=List_{request_code}.xlsx"})


@router.get("/export-csv-all")
def export_all_operators_to_csv_stream(ids: str = None, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    user_role_str = get_user_role_str(current_user)
    
    query = db.query(ReactivationOperator, OperatorReactivationRequest, District.district_name).\
        join(OperatorReactivationRequest, ReactivationOperator.request_id == OperatorReactivationRequest.id).\
        outerjoin(District, OperatorReactivationRequest.district_id == District.district_code)
    
    if ids:
        id_tokens = [code.strip() for code in ids.split(",") if code.strip()]
        if id_tokens:
            if all(token.isdigit() for token in id_tokens):
                operator_ids = [int(token) for token in id_tokens]
                query = query.filter(ReactivationOperator.id.in_(operator_ids))
            else:
                query = query.filter(OperatorReactivationRequest.request_code.in_(id_tokens))
    else:
        query = query.filter(ReactivationOperator.status_id.notin_([StatusEnum.REVERTED.value, StatusEnum.APPROVED.value, StatusEnum.REJECTED.value]))
        

    if user_role_str == "dc":
        query = query.filter(OperatorReactivationRequest.district_id == str(current_user.district_id))
        
    results = query.all()
    
    export_data = []
    for i, (o, req, dist_name) in enumerate(results):
        status_upper = (o.status or '').upper().strip()
        if status_upper == 'PENDING':
            reviewed_time = ""
        else:
            reviewed_time = req.updated_at.strftime("%Y-%m-%d %H:%M:%S") if req.updated_at else "—"
        
        display_status = "APPROVED" if status_upper in ["ACTIVE", "APPROVED"] else status_upper

        export_data.append({
            "s_no": i + 1,
            "request_code": req.request_code,
            "district_name": dist_name,
            "role": o.role,
            "operator_name": o.operator_name,
            "aadhaar_number": o.aadhaar_number,
            "registrar_code": o.registrar_code,
            "ea_code": o.ea_code,
            "user_code": o.user_code,
            "certificate_number": o.certificate_number,
            "lms_certificate_id": o.lms_certificate_id,
            "operator_mobile": o.operator_mobile,
            "email_id": o.email_id,
            "status": display_status,
            "submitted_at": req.created_at.strftime("%Y-%m-%d %H:%M:%S") if req.created_at else "—",
            "reviewed_at": reviewed_time
        })
        
    column_mappings = {
        "s_no": "S.No", 
        "request_code": "Batch Request ID", 
        "district_name": "District Name",
        "role": "Role Profile", 
        "operator_name": "Name As Per Aadhaar", 
        "aadhaar_number": "Aadhaar Number",
        "registrar_code": "Registrar Code",
        "ea_code": "EA Code", 
        "user_code": "User Code", 
        "certificate_number": "NSEIT Certificate Number",
        "lms_certificate_id": "LMS Certificate ID", 
        "operator_mobile": "Mobile Phone Number",
        "email_id": "Primary E-MAIL ID", 
        "status": "Current Audit Status", 
        "submitted_at": "Submission Date",
        "reviewed_at": "Reviewed Time"
    }
    return generate_csv_export(export_data, column_mappings, "All_Pending_Reactivation_Operators")


@router.get("/export-csv-uidai")
def export_uidai_operators_to_csv_stream(ids: str = None, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    user_role_str = get_user_role_str(current_user)
    
    query = db.query(ReactivationOperator, OperatorReactivationRequest, District.district_name).\
        join(OperatorReactivationRequest, ReactivationOperator.request_id == OperatorReactivationRequest.id).\
        outerjoin(District, OperatorReactivationRequest.district_id == District.district_code)
    
    if ids:
        id_tokens = [code.strip() for code in ids.split(",") if code.strip()]
        if id_tokens:
            if all(token.isdigit() for token in id_tokens):
                operator_ids = [int(token) for token in id_tokens]
                query = query.filter(ReactivationOperator.id.in_(operator_ids))
            else:
                query = query.filter(OperatorReactivationRequest.request_code.in_(id_tokens))
    else:
        query = query.filter(ReactivationOperator.status_id == StatusEnum.SENT_TO_UIDAI.value)
        

    if user_role_str == "dc":
        query = query.filter(OperatorReactivationRequest.district_id == str(current_user.district_id))
        
    results = query.all()
    
    export_data = []
    for i, (o, req, dist_name) in enumerate(results):
        export_data.append({
            "s_no": i + 1,
            "role": o.role or "Supervisor",
            "operator_name": o.operator_name or "—",
            "registrar_code": o.registrar_code or "986",
            "ea_code": o.ea_code or "—",
            "user_code": o.user_code or "—",
            "certificate_number": o.certificate_number or "—",
            "operator_mobile": o.operator_mobile or "—",
            "email_id": o.email_id or "—",
            "aadhaar_number": f"XXXX-XXXX-{o.aadhaar_number[-4:]}" if o.aadhaar_number and len(o.aadhaar_number) >= 4 else (o.aadhaar_number or "—"),
            "certification_date": str(o.certification_date) if o.certification_date else "—",
            "remarks": "",
            "model": o.model_type or "—"
        })
        
    column_mappings = {
        "s_no": "S.No", 
        "role": "Role", 
        "operator_name": "Name as per Aadhaar", 
        "registrar_code": "Registrar Code",
        "ea_code": "EA Code", 
        "user_code": "User Code", 
        "certificate_number": "Certificate Number",
        "operator_mobile": "Mobile",
        "email_id": "Primary E-MAIL ID", 
        "aadhaar_number": "Aadhaar Number",
        "certification_date": "Certification Date",
        "remarks": "Remarks",
        "model": "Model"
    }
    return generate_csv_export(export_data, column_mappings, "UIDAI_Sent_Reactivation_Operators")

@router.get("/requests/{request_code}/files/{file_type}")
async def get_reactivation_file(request_code: str, file_type: str, db: Session = Depends(get_db)):
    req = db.query(OperatorReactivationRequest).filter(OperatorReactivationRequest.request_code == request_code).first()
    if not req:
        raise HTTPException(status_code=404, detail="Request not found")
        
    doc = db.query(ReactivationDocument).filter(
        ReactivationDocument.request_id == req.id,

        ReactivationDocument.doc_type == file_type
    ).first()
    if not doc or not os.path.exists(doc.path):
        raise HTTPException(status_code=404, detail="File not found")
    
    mime_type, _ = mimetypes.guess_type(doc.path)
    if not mime_type:
        mime_type = "application/octet-stream"
        
    return FileResponse(
        path=doc.path,
        filename=doc.original_filename,
        media_type=mime_type,
        content_disposition_type="inline"
    )

@router.get("/search-suspended-operators")
async def search_suspended_operators(q: str = "", db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    user_role_str = get_user_role_str(current_user)
    
    import re
    def norm_mob(m: str | None) -> str:
        if not m:
            return ""
        digits = re.sub(r"\D", "", str(m))
        return digits[-10:] if len(digits) >= 10 else digits

    def norm_str(s: str | None) -> str:
        if not s:
            return ""
        return str(s).strip().lower()

    # Query all operators who have already applied in ReactivationOperator (irrespective of status)
    applied_records = db.query(
        ReactivationOperator.operator_mobile,
        ReactivationOperator.email_id,
        ReactivationOperator.user_code,
        ReactivationOperator.certificate_number
    ).all()

    applied_mobiles = {norm_mob(r[0]) for r in applied_records if r[0]}
    applied_emails = {norm_str(r[1]) for r in applied_records if r[1]}
    applied_user_codes = {norm_str(r[2]) for r in applied_records if r[2]}
    applied_certs = {norm_str(r[3]) for r in applied_records if r[3]}

    applied_mobiles.discard("")
    applied_emails.discard("")
    applied_user_codes.discard("")
    applied_certs.discard("")

    from sqlalchemy import or_
    query = db.query(Operator).filter(
        or_(
            Operator.status.ilike("%suspended%"),
            Operator.status.ilike("%inactive%"),
            Operator.status.ilike("%deactive%"),
            Operator.status.ilike("%deboard%"),
            Operator.inactive_reason.ilike("%suspended%")
        )
    )
    if user_role_str == "dc":
        query = query.filter(
            or_(
                Operator.mapped_dc_id == current_user.id,
                Operator.district_id == str(current_user.district_id),
                Operator.district_id == current_user.district_id
            )
        )
        
    if q and q.strip():
        search_term = f"%{q.strip().lower()}%"
        query = query.filter(
            (Operator.name.ilike(search_term)) |
            (Operator.mobile.ilike(search_term)) |
            (Operator.email.ilike(search_term)) |
            (Operator.user_code.ilike(search_term)) |
            (Operator.nseit_certificate_number.ilike(search_term))
        )
        
    operators = query.order_by(Operator.id.desc()).limit(50).all()

    # Fallback: if no operators found for DC filter with search term, search across all operators
    if not operators and q and q.strip():
        search_term = f"%{q.strip().lower()}%"
        fallback_query = db.query(Operator).filter(
            (Operator.name.ilike(search_term)) |
            (Operator.mobile.ilike(search_term)) |
            (Operator.email.ilike(search_term)) |
            (Operator.user_code.ilike(search_term)) |
            (Operator.nseit_certificate_number.ilike(search_term))
        )
        operators = fallback_query.order_by(Operator.id.desc()).limit(50).all()

    # Fallback: if empty query, return top 50 operators
    if not operators and not (q and q.strip()):
        operators = db.query(Operator).order_by(Operator.id.desc()).limit(50).all()

    results = []
    seen_ids = set()
    for o in operators:
        mob = norm_mob(o.mobile)
        em = norm_str(o.email)
        uc = norm_str(o.user_code)
        cert = norm_str(o.nseit_certificate_number)

        # Exclude operators who have already applied at least once (irrespective of current status)
        if (mob and mob in applied_mobiles) or \
           (em and em in applied_emails) or \
           (uc and uc in applied_user_codes) or \
           (cert and cert in applied_certs):
            continue

        if o.id in seen_ids:
            continue
        seen_ids.add(o.id)

        clean_mob_val = str(o.mobile).strip() if o.mobile else ""
        if clean_mob_val.endswith(".0"):
            clean_mob_val = clean_mob_val[:-2]

        results.append({
            "id": o.id,
            "name": o.name,
            "mobile": clean_mob_val if clean_mob_val and clean_mob_val != "None" else "—",
            "email": o.email if o.email and o.email != "None" else "—",
            "role": o.role or "Operator",
            "nseit_id": o.nseit_certificate_number or "",
            "user_code": o.user_code or f"OP-{o.id:04d}",
            "registrar_code": o.registrar_code or "986",
            "ea_code": o.ea_code or "2084",
            "aadhaar_last4": o.aadhaar_last4 or ""
        })

    return results


class ExportAndMailReactivationRequest(BaseModel):
    ids: str | None = None
    email_to: str | None = None
    email_cc: str | None = None
    email_bcc: str | None = None
    subject: str | None = None
    body_html: str | None = None
    attach_csv: bool = True
    custom_files: list[dict] | None = None

@router.get("/export-and-mail/recipient")
def get_reactivation_export_mail_recipient():
    from backend.utils.email_utils import DEFAULT_UIDAI_RECIPIENT_EMAIL
    return {"recipient_email": DEFAULT_UIDAI_RECIPIENT_EMAIL}

@router.post("/export-and-mail")
def export_and_mail_reactivation_to_uidai(
    payload: ExportAndMailReactivationRequest, 
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    import csv
    import io
    import asyncio
    from backend.utils.email_utils import send_uidai_export_email, DEFAULT_UIDAI_RECIPIENT_EMAIL

    recipient = payload.email_to.strip() if payload.email_to else DEFAULT_UIDAI_RECIPIENT_EMAIL

    query = db.query(ReactivationOperator, OperatorReactivationRequest, District.district_name)\
              .join(OperatorReactivationRequest, ReactivationOperator.request_id == OperatorReactivationRequest.id)\
              .outerjoin(District, OperatorReactivationRequest.district_id == District.district_code)

    if payload.ids:
        id_list = [int(i.strip()) for i in payload.ids.split(",") if i.strip().isdigit()]
        query = query.filter(
            or_(
                OperatorReactivationRequest.id.in_(id_list),
                ReactivationOperator.id.in_(id_list)
            )
        )
    else:
        query = query.filter(
            OperatorReactivationRequest.status_id.in_([
                StatusEnum.PENDING.value,
                StatusEnum.REAPPLIED.value
            ]),
            or_(OperatorReactivationRequest.is_mailed == 0, OperatorReactivationRequest.is_mailed.is_(None))
        )

    records = query.order_by(OperatorReactivationRequest.created_at.desc()).all()

    if not records:
        raise HTTPException(status_code=400, detail="No reactivation requests found matching the selection.")

    stream = io.StringIO()
    writer = csv.writer(stream)

    headers = [
        "S.No",
        "Role",
        "Name as per Aadhaar",
        "Registrar Code",
        "EA Code",
        "User Code",
        "Certificate Number",
        "Mobile",
        "Primary E-MAIL ID",
        "Aadhaar Number",
        "Certification Date",
        "Remarks",
        "Model"
    ]
    writer.writerow(headers)

    for idx, (op, req, dist_name) in enumerate(records, start=1):
        writer.writerow([
            idx,
            op.role or "Supervisor",
            op.operator_name or "—",
            op.registrar_code or "986",
            op.ea_code or "—",
            op.user_code or "—",
            op.certificate_number or "—",
            op.operator_mobile or "—",
            op.email_id or "—",
            f"XXXX-XXXX-{op.aadhaar_number[-4:]}" if op.aadhaar_number and len(op.aadhaar_number) >= 4 else (op.aadhaar_number or "—"),
            str(op.certification_date) if op.certification_date else "—",
            "",
            op.model_type or "—"
        ])

    csv_content = stream.getvalue()
    target_email = (payload.email_to or DEFAULT_UIDAI_RECIPIENT_EMAIL).strip()

    try:
        asyncio.run(send_uidai_export_email(
            csv_content=csv_content,
            record_count=len(records),
            module_name="Operator Reactivation",
            filename="operator_reactivation_sent_to_uidai.csv",
            email_to=target_email,
            email_cc=payload.email_cc,
            email_bcc=payload.email_bcc,
            custom_subject=payload.subject,
            custom_body_html=payload.body_html,
            attach_csv=payload.attach_csv,
            custom_files=payload.custom_files
        ))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to email CSV export: {str(e)}")

    now_ist = get_ist_now()
    reviewer_id = getattr(current_user, 'id', 1)
    processed_requests = set()
    for (op, req, dist_name) in records:
        if req not in processed_requests:
            req.is_mailed = 1
            req.updated_at = now_ist
            req.reviewed_by = reviewer_id
            processed_requests.add(req)

    db.commit()

    return {
        "success": True,
        "detail": f"Export CSV ({len(records)} records) emailed successfully to {target_email} and moved to Under Processing queue.",
        "recipient_email": target_email,
        "count": len(records)
    }



